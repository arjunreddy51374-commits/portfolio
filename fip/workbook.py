"""Create a native Excel report from one approved, immutable FIP release."""
from __future__ import annotations

from io import BytesIO
import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Side, Border
from openpyxl.utils import get_column_letter


NAVY = "15324B"
TEAL = "197C78"
PALE = "EAF2F5"
INK = "243746"
MUTED = "637887"
WHITE = "FFFFFF"


def _text(value: Any) -> Any:
    """Preserve report values as text, including exact decimal strings."""
    if isinstance(value, str):
        return value
    if value is None:
        return ""
    return str(value)


def _set_literal(cell, value: Any) -> None:
    cell.value = _text(value)
    if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@", "\t", "\r")):
        cell.data_type = "s"


def _table(workbook: Workbook, title: str, headers: list[str], rows: list[list[Any]]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.sheet_view.showGridLines = False
    sheet.append(headers)
    for row in rows:
        sheet.append([_text(value) for value in row])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    header_fill = PatternFill("solid", fgColor=NAVY)
    header_font = Font(name="Arial", size=10, bold=True, color=WHITE)
    bottom = Side(style="thin", color=WHITE)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=bottom)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color=INK)
            cell.alignment = Alignment(vertical="top", wrap_text=False)
            if cell.value and isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@", "\t", "\r")):
                cell.data_type = "s"
    sheet.row_dimensions[1].height = 30
    for index, header in enumerate(headers, 1):
        values = [len(str(header))]
        values.extend(min(len(str(sheet.cell(r, index).value or "")), 56) for r in range(2, sheet.max_row + 1))
        sheet.column_dimensions[get_column_letter(index)].width = min(max(max(values, default=12) + 2, 12), 38)
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0


def build_release_workbook(publication: dict[str, Any]) -> BytesIO:
    release = publication["release"]
    observations = publication.get("observations", [])
    calculations = publication.get("calculations", [])
    fx_rates = publication.get("fx_rates", [])

    workbook = Workbook()
    dashboard = workbook.active
    dashboard.title = "Dashboard"
    dashboard.sheet_view.showGridLines = False
    dashboard.freeze_panes = "A5"
    dashboard["B2"] = "Financial Intelligence Platform"
    dashboard["B2"].font = Font(name="Arial", size=16, bold=True, color=NAVY)
    dashboard["B3"] = f"Approved release {release['version']}"
    dashboard["B3"].font = Font(name="Arial", size=11, color=MUTED)
    dashboard["B5"] = "Release details"
    dashboard["B5"].font = Font(name="Arial", size=11, bold=True, color=WHITE)
    dashboard["B5"].fill = PatternFill("solid", fgColor=TEAL)
    details = [
        ("Release ID", release["id"]), ("Release version", release["version"]),
        ("Approved at (UTC)", release.get("approved_at", "")),
        ("Release notes", release.get("notes", "")),
        ("Reviewer decision", release.get("approval_reason", "")),
        ("Formula catalogue", release.get("formula_catalog_version", "1.0.0")),
        ("Published observations", len(observations)),
        ("Published calculation runs", len(calculations)),
        ("Entities represented", len({r.get("entity_id") for r in observations + calculations})),
        ("Jurisdictions", ", ".join(sorted({r.get("jurisdiction_code", "") for r in observations if r.get("jurisdiction_code")}))),
        ("Reporting periods", ", ".join(sorted({r.get("period_id", "") for r in observations + calculations if r.get("period_id")}))),
        ("Workbook purpose", "Read-only view of this approved FIP release snapshot."),
        ("Calculation policy", "Values and calculated results are copied from the canonical FIP engine; this workbook does not recalculate them."),
    ]
    for row_number, (label, value) in enumerate(details, 6):
        dashboard.cell(row_number, 2, label).font = Font(name="Arial", size=10, bold=True, color=INK)
        _set_literal(dashboard.cell(row_number, 3), value)
        dashboard.cell(row_number, 3).font = Font(name="Arial", size=10, color=INK)
        dashboard.cell(row_number, 2).fill = PatternFill("solid", fgColor=PALE)
        dashboard.cell(row_number, 2).alignment = Alignment(vertical="top", wrap_text=True)
        dashboard.cell(row_number, 3).alignment = Alignment(vertical="top", wrap_text=True)
    dashboard.column_dimensions["A"].width = 3
    dashboard.column_dimensions["B"].width = 30
    dashboard.column_dimensions["C"].width = 100
    dashboard.sheet_properties.pageSetUpPr.fitToPage = True

    raw_headers = ["Observation ID", "Entity", "Jurisdiction", "Domain", "Metric ID", "Reported label",
                   "Reported value (exact decimal)", "Currency", "Unit", "Period ID", "Period start", "Period end",
                   "Period type", "Reporting basis", "Framework", "Framework version", "Fiscal scope", "Budget stage",
                   "Source title", "Source URL", "Publication date", "Evidence location", "Freshness", "Status", "Version",
                   "Supersedes observation", "FX rate ID", "FX convention", "Validation message", "Validation checks JSON"]
    raw_rows = [[item.get(key, "") for key in (
        "id", "canonical_name", "jurisdiction_code", "domain", "metric_id", "reported_label", "value_decimal", "currency", "unit",
        "period_id", "period_start", "period_end", "period_type", "reporting_basis", "framework_id", "framework_version", "fiscal_scope",
        "budget_stage", "source_title", "source_url", "publication_date", "evidence_reference", "freshness_status", "status", "version",
        "supersedes_id", "fx_rate_id", "conversion_convention", "validation_message", "validation_json")]
        for item in observations]
    _table(workbook, "Raw Data", raw_headers, raw_rows)

    standardized_headers = ["Observation ID", "Entity", "Domain", "Metric ID", "Original amount (exact decimal)",
                            "Original currency", "Original unit", "Standardized amount (exact decimal)", "Standardized currency",
                            "FX rate ID", "FX rate", "Conversion convention", "Period ID", "Framework", "Framework version",
                            "Source title", "Evidence location"]
    rates_by_id = {rate["id"]: rate for rate in fx_rates}
    standardized_rows = []
    for item in observations:
        rate = rates_by_id.get(item.get("fx_rate_id"), {})
        standardized_rows.append([item.get("id"), item.get("canonical_name"), item.get("domain"), item.get("metric_id"),
                                  item.get("value_decimal"), item.get("currency"), item.get("unit"),
                                  item.get("standardized_value_decimal"), item.get("standardized_currency"),
                                  item.get("fx_rate_id"), rate.get("rate_decimal", ""), item.get("conversion_convention"),
                                  item.get("period_id"), item.get("framework_id"), item.get("framework_version"),
                                  item.get("source_title"), item.get("evidence_reference")])
    _table(workbook, "Standardized Data", standardized_headers, standardized_rows)

    statement_headers = ["Entity", "Jurisdiction", "Domain", "Metric", "Reported label", "Reported amount (exact decimal)",
                         "Original unit", "Original currency", "Standardized amount (exact decimal)", "Standardized currency",
                         "Period", "Period start", "Period end", "Reporting basis", "Framework", "Framework version",
                         "Fiscal scope", "Budget stage", "Source", "Evidence location", "Observation version"]
    statement_rows = [[item.get("canonical_name"), item.get("jurisdiction_code"), item.get("domain"), item.get("metric_id"),
                       item.get("reported_label"), item.get("value_decimal"), item.get("unit"), item.get("currency"),
                       item.get("standardized_value_decimal"), item.get("standardized_currency"), item.get("period_id"),
                       item.get("period_start"), item.get("period_end"), item.get("reporting_basis"), item.get("framework_id"),
                       item.get("framework_version"), item.get("fiscal_scope"), item.get("budget_stage"), item.get("source_title"),
                       item.get("evidence_reference"), item.get("version")] for item in observations]
    _table(workbook, "Statements", statement_headers, statement_rows)

    engine_headers = ["Calculation ID", "Entity", "Domain", "Period", "Metric ID", "Metric", "Value (exact decimal)", "Unit",
                      "Formula ID", "Formula version", "Formula", "Inputs JSON", "Input observations JSON",
                      "Framework", "Framework version", "Status", "Warnings and limitations"]
    engine_rows = []
    for run in calculations:
        result = run.get("result", {})
        input_refs = result.get("input_observations", [])
        for metric_id, metric in result.get("metrics", {}).items():
            engine_rows.append([run.get("id"), run.get("canonical_name"), run.get("domain"), run.get("period_id"), metric_id,
                                metric.get("label"), metric.get("value_decimal"), metric.get("unit"), metric.get("formula_id"),
                                metric.get("formula_version"), metric.get("formula"), _json_text(metric.get("inputs", {})),
                                _json_text(input_refs), result.get("framework_id"), result.get("framework_version"),
                                run.get("status"), "; ".join(result.get("limitations", []))])
    _table(workbook, "Engine", engine_headers, engine_rows)

    corporate_headers = ["Entity", "Period", "Metric", "Value (exact decimal)", "Currency / unit", "Formula", "Source fact IDs"]
    corp_rows = _domain_rows(observations, calculations, "corporate")
    _table(workbook, "Corporate", corporate_headers, corp_rows)
    gov_rows = _domain_rows(observations, calculations, "government")
    _table(workbook, "Government", corporate_headers, gov_rows)

    sources: dict[str, dict[str, Any]] = {}
    for item in observations:
        if item.get("source_id"):
            sources[item["source_id"]] = {"Source ID": item.get("source_id"), "Title": item.get("source_title"),
                                          "Publisher": item.get("publisher", ""), "Category": item.get("source_category", ""),
                                          "Jurisdiction": item.get("source_jurisdiction", ""), "URL": item.get("source_url"),
                                          "Publication date": item.get("publication_date"), "Retrieval method": item.get("retrieval_method", ""),
                                          "Expected frequency": item.get("expected_frequency", ""), "Licence": item.get("license_name", ""),
                                          "Reuse status": item.get("permitted_use", "public_display")}
    for rate in fx_rates:
        if rate.get("source_id"):
            sources.setdefault(rate["source_id"], {"Source ID": rate.get("source_id"), "Title": rate.get("source_title"),
                                                    "Publisher": rate.get("publisher", ""), "Category": rate.get("source_category", ""),
                                                    "Jurisdiction": rate.get("source_jurisdiction", ""), "URL": rate.get("source_url"),
                                                    "Publication date": rate.get("publication_date", ""), "Retrieval method": rate.get("retrieval_method", ""),
                                                    "Expected frequency": rate.get("expected_frequency", ""), "Licence": rate.get("license_name", ""),
                                                    "Reuse status": rate.get("permitted_use", "public_display")})
    source_headers = ["Source ID", "Title", "Publisher", "Category", "Jurisdiction", "URL", "Publication date",
                      "Retrieval method", "Expected frequency", "Licence", "Reuse status"]
    _table(workbook, "Source Register", source_headers, [[source.get(header, "") for header in source_headers] for source in sources.values()])

    fx_headers = ["FX Rate ID", "Source currency", "Target currency", "Rate (exact decimal)", "Period ID", "Convention", "Source", "Source URL"]
    fx_rows = [[rate.get("id"), rate.get("source_currency"), rate.get("target_currency"), rate.get("rate_decimal"),
                rate.get("period_id"), rate.get("conversion_convention"), rate.get("source_title"), rate.get("source_url")] for rate in fx_rates]
    _table(workbook, "FX Rates", fx_headers, fx_rows)

    audit_headers = ["Record type", "Record ID", "Action", "Actor", "Decision date", "Decision reason / release note"]
    audit_rows = [["Publication release", release.get("id"), "RELEASE_APPROVED", "reviewer", release.get("approved_at"),
                   release.get("approval_reason", "")],
                  ["Publication release", release.get("id"), "RELEASE_SNAPSHOT_CREATED", release.get("created_by", "admin"), release.get("created_at"),
                   release.get("notes", "")]]
    _table(workbook, "Audit", audit_headers, audit_rows)

    controls_headers = ["Control", "Value", "Description"]
    controls_rows = [
        ["Release ID", release.get("id"), "Immutable publication snapshot ID."],
        ["Release version", release.get("version"), "Version used by Web, Excel and Power BI delivery."],
        ["Formula catalogue version", release.get("formula_catalog_version", "1.0.0"), "Formula definitions used by the canonical calculation service."],
        ["Data entry", "Manual transcription", "This prototype does not retrieve or extract source documents."],
        ["Numeric storage", "Exact decimal strings", "Displayed amounts are preserved as text to avoid spreadsheet precision loss."],
        ["Workbook calculations", "None", "Calculated metrics are exported from the canonical engine with inputs and formulas."],
        ["Comparability", "Framework, basis, period and fiscal scope", "Review these fields before comparing records."],
        ["Rights", "Approved public-display records only", "Sources with link-only or internal-only rights are excluded."],
        ["Audit scope", "Release metadata", "Detailed staff-only lifecycle events remain in the authenticated audit API."],
    ]
    _table(workbook, "Controls", controls_headers, controls_rows)

    desired_order = ["Dashboard", "Controls", "Corporate", "Government", "Statements", "Engine",
                     "Raw Data", "Standardized Data", "Source Register", "FX Rates", "Audit"]
    for target, title in enumerate(desired_order):
        sheet = workbook[title]
        current = workbook.worksheets.index(sheet)
        workbook.move_sheet(sheet, offset=target - current)
    workbook.active = workbook.worksheets.index(dashboard)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _json_text(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def _domain_rows(observations: list[dict[str, Any]], calculations: list[dict[str, Any]], domain: str) -> list[list[Any]]:
    rows = []
    for item in observations:
        if item.get("domain") == domain:
            rows.append([item.get("canonical_name"), item.get("period_id"), item.get("metric_id"),
                         item.get("standardized_value_decimal"), f"{item.get('standardized_currency', '')} normalized units",
                         "Reported fact", item.get("id")])
    for run in calculations:
        if run.get("domain") != domain:
            continue
        result = run.get("result", {})
        refs = ", ".join(item.get("id", "") for item in result.get("input_observations", []))
        for metric_id, metric in result.get("metrics", {}).items():
            rows.append([run.get("canonical_name"), run.get("period_id"), metric_id, metric.get("value_decimal"),
                         metric.get("unit"), metric.get("formula"), refs])
    return rows

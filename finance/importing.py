"""Strict in-memory CSV/XLSX statement-history import."""
from __future__ import annotations

import csv
import io
import zipfile
from collections import OrderedDict
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from finance.calculations import AMOUNT_FIELDS, FinancialInputError
from finance.history import analyze_history


MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_ROWS = 500
MAX_COMPANIES = 20
MAX_PERIODS_PER_COMPANY = 20
MAX_XLSX_UNCOMPRESSED_BYTES = 25 * 1024 * 1024
REQUIRED_COLUMNS = {"company_name", "period", "currency", "unit", "source_url"}
METADATA_COLUMNS = REQUIRED_COLUMNS | {"publication_date", "prior_period"}
ALLOWED_COLUMNS = METADATA_COLUMNS | AMOUNT_FIELDS
NOT_APPLICABLE_MARKERS = {"n/a", "na", "not applicable", "not_applicable"}


def _header_key(value: Any) -> str:
    key = str(value or "").strip().lower().replace("-", "_").replace(" ", "_")
    while "__" in key:
        key = key.replace("__", "_")
    if key in {"company", "issuer", "entity"}:
        return "company_name"
    if key in {"fiscal_period", "reporting_period", "year"}:
        return "period"
    return key


def _string(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float):
        return format(Decimal(str(value)), "f")
    return str(value).strip()


def _records(headers: list[Any], rows: list[list[Any]]) -> list[tuple[int, dict[str, Any]]]:
    keys = [_header_key(value) for value in headers]
    if not keys or any(not key for key in keys):
        raise FinancialInputError("The import file must have a non-empty header row.")
    if len(set(keys)) != len(keys):
        raise FinancialInputError("The import file contains duplicate column headers.")
    unknown = sorted(set(keys) - ALLOWED_COLUMNS)
    if unknown:
        raise FinancialInputError("Unknown import columns: " + ", ".join(unknown))
    missing = sorted(REQUIRED_COLUMNS - set(keys))
    if missing:
        raise FinancialInputError("Missing required import columns: " + ", ".join(missing))
    output = []
    for row_number, row in enumerate(rows, start=2):
        if not any(value not in (None, "") for value in row):
            continue
        if len(row) != len(keys):
            raise FinancialInputError(f"Row {row_number}: the number of values does not match the header columns.")
        if len(output) >= MAX_ROWS:
            raise FinancialInputError(f"Import is limited to {MAX_ROWS} non-empty data rows.")
        values = {keys[index]: (row[index] if index < len(row) else None) for index in range(len(keys))}
        output.append((row_number, values))
    if not output:
        raise FinancialInputError("The import file contains no financial data rows.")
    return output


def _csv_records(content: bytes) -> list[tuple[int, dict[str, Any]]]:
    try:
        text = content.decode("utf-8-sig")
        reader = csv.reader(io.StringIO(text, newline=""), strict=True)
        rows = list(reader)
    except (UnicodeDecodeError, csv.Error):
        raise FinancialInputError("CSV files must be valid UTF-8 with consistent comma-separated rows.") from None
    if not rows:
        raise FinancialInputError("The CSV file is empty.")
    return _records(rows[0], rows[1:])


def _xlsx_records(content: bytes) -> list[tuple[int, dict[str, Any]]]:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 100 or sum(item.file_size for item in entries) > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise FinancialInputError("The workbook exceeds safe archive size limits.")
            if not any(item.filename == "xl/workbook.xml" for item in entries):
                raise FinancialInputError("The uploaded file is not a valid XLSX workbook.")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except FinancialInputError:
        raise
    except Exception:
        raise FinancialInputError("The uploaded file is not a readable XLSX workbook.") from None
    try:
        if not workbook.worksheets:
            raise FinancialInputError("The workbook has no worksheets.")
        sheet = workbook.worksheets[0]
        if (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) > len(ALLOWED_COLUMNS):
            raise FinancialInputError("The first worksheet exceeds the supported row or column limit.")
        iterator = sheet.iter_rows(values_only=False)
        first_row = next(iterator, None)
        if first_row is None:
            raise FinancialInputError("The workbook is empty.")
        headers = [cell.value for cell in first_row]
        rows = []
        for row in iterator:
            values = []
            for cell in row:
                if cell.data_type == "f":
                    raise FinancialInputError("Formula cells are not accepted. Import evaluated statement values instead.")
                values.append(cell.value)
            rows.append(values)
            if len(rows) > MAX_ROWS:
                raise FinancialInputError(f"Import is limited to {MAX_ROWS} data rows.")
        return _records(headers, rows)
    finally:
        workbook.close()


def parse_statement_upload(filename: str, content: bytes) -> list[dict[str, Any]]:
    if not isinstance(filename, str) or not filename:
        raise FinancialInputError("Choose a CSV or XLSX file.")
    if len(content) > MAX_UPLOAD_BYTES:
        raise FinancialInputError("Import files are limited to 5 MiB.")
    extension = Path(filename).suffix.lower()
    if extension == ".csv":
        records = _csv_records(content)
    elif extension == ".xlsx":
        records = _xlsx_records(content)
    else:
        raise FinancialInputError("Only .csv and .xlsx financial statement files are supported.")

    groups: OrderedDict[str, dict[str, Any]] = OrderedDict()
    for row_number, values in records:
        company = _string(values.get("company_name"))
        if not company or len(company) > 200:
            raise FinancialInputError(f"Row {row_number}: company_name must contain 1 to 200 characters.")
        group_key = company.casefold()
        if group_key not in groups:
            if len(groups) >= MAX_COMPANIES:
                raise FinancialInputError(f"Import is limited to {MAX_COMPANIES} companies.")
            groups[group_key] = {"company_name": company, "periods": []}
        if len(groups[group_key]["periods"]) >= MAX_PERIODS_PER_COMPANY:
            raise FinancialInputError(f"Row {row_number}: each company is limited to {MAX_PERIODS_PER_COMPANY} periods.")

        statement = {}
        for key in ("period", "currency", "unit", "source_url", "publication_date", "prior_period"):
            value = values.get(key)
            if value not in (None, ""):
                statement[key] = _string(value)
        statement["currency"] = statement.get("currency", "").upper()
        statement["source_kind"] = "PUBLIC_SOURCE"
        amounts = {}
        not_applicable = []
        for field in AMOUNT_FIELDS:
            raw = values.get(field)
            if raw in (None, ""):
                continue
            if isinstance(raw, str) and raw.strip().lower() in NOT_APPLICABLE_MARKERS:
                not_applicable.append(field)
                continue
            amounts[field] = _string(raw)
        statement["amounts"] = amounts
        statement["not_applicable_fields"] = not_applicable
        groups[group_key]["periods"].append(statement)

    return list(groups.values())


def analyze_statement_upload(filename: str, content: bytes) -> dict[str, Any]:
    companies = [analyze_history(group) for group in parse_statement_upload(filename, content)]
    currencies = {company["currency"] for company in companies}
    peer_comparison = None
    warnings = []
    if len(companies) >= 2:
        if len(currencies) != 1:
            warnings.append("Peer comparison is withheld because the companies use different currencies and no reviewed FX conversion was supplied.")
        else:
            peer_comparison = []
            for company in companies:
                latest = company["periods"][-1]
                peer_comparison.append({
                    "company_name": company["company_name"],
                    "period": latest["metadata"]["period"],
                    "currency": company["currency"],
                    "source_unit": latest["metadata"]["unit"],
                    "calculation_unit": "base currency units",
                    "normalized_inputs": latest["normalized_inputs"],
                    "metrics": latest["metrics"],
                    "source_url": latest["metadata"]["source_url"],
                })
    return {
        "companies": companies,
        "peer_comparison": peer_comparison,
        "warnings": warnings,
        "row_count": sum(len(company["periods"]) for company in companies),
        "stored": False,
    }
